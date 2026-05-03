# pitch_deck.py
# Generates a 5-slide client pitch deck for any Sleeping Giant.

import io
import json

from dotenv import load_dotenv
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

from local_fallbacks import fallback_pitch_content
from llm_client import generate_text

load_dotenv()

ORANGE = RGBColor(0xFF, 0x45, 0x00)
DARK = RGBColor(0x1A, 0x1A, 0x2E)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
RED = RGBColor(0xDC, 0x35, 0x45)
GREEN = RGBColor(0x28, 0xA7, 0x45)
YELLOW = RGBColor(0xFF, 0xC1, 0x07)


def _extract_json(raw):
    raw = raw.strip()
    if "```" in raw:
        raw = raw.split("```")[1]
        if raw.startswith("json"):
            raw = raw[4:]
    return json.loads(raw.strip())


def generate_pitch_content(giant):
    """Ask Claude to write all 5 slides of pitch deck content."""
    weaknesses = "\n".join([f"- {w}" for w in giant.get("weaknesses", [])])
    bullets_original = "\n".join([f"- {b}" for b in giant.get("bullets", [])])
    rewrite = giant.get("rewrite", "Rewrite not available")
    revenue = giant.get("estimated_revenue", 0)
    money_left = giant.get("money_left_on_table", 0)
    uplift = giant.get("uplift_percent", 0)
    score = giant.get("listing_score", 0)

    prompt = f"""
You are a senior Amazon marketing consultant creating a client pitch deck.
The goal: convince this Amazon seller to hire us to fix their listing.

PRODUCT DATA:
Title: {giant.get('title', 'Unknown Product')}
Reviews: {giant.get('review_count', 0)} reviews at {giant.get('rating', 0)} stars
Listing Quality Score: {score}/100
Estimated Monthly Revenue: Rs.{revenue:,}
Revenue Being Lost: Rs.{money_left:,}/month
Potential Uplift: {uplift:.0f}%

LISTING WEAKNESSES:
{weaknesses}

ORIGINAL BULLETS:
{bullets_original}

AI REWRITE PREVIEW:
{rewrite[:600]}

Write content for exactly 5 slides. Be punchy, specific, and persuasive.
Use their actual numbers. No generic filler.

Respond only in this exact JSON format:

{{
  "slide1": {{
    "title": "Product Snapshot",
    "product_name": "<shortened product name under 50 chars>",
    "monthly_revenue": "Rs.{revenue:,}/mo",
    "star_rating": "{giant.get('rating', 0)} stars",
    "review_count": "{giant.get('review_count', 0):,} reviews",
    "listing_score": "{score}/100",
    "tagline": "<one punchy line about this product's situation>"
  }},
  "slide2": {{
    "title": "What's Hurting Your Sales",
    "subtitle": "<one line framing the problem>",
    "weakness_1": "<specific weakness with impact>",
    "weakness_2": "<specific weakness with impact>",
    "weakness_3": "<specific weakness with impact>",
    "closing_line": "<one alarming but true statement about what bad listings cost>"
  }},
  "slide3": {{
    "title": "The Revenue Opportunity",
    "current_revenue": "Rs.{revenue:,}",
    "optimized_revenue": "Rs.{giant.get('optimized_revenue', revenue):,}",
    "money_left": "Rs.{money_left:,}",
    "uplift_percent": "{uplift:.0f}%",
    "stat_1": "<compelling industry stat about listing optimization>",
    "stat_2": "<second compelling stat>",
    "insight": "<one sentence insight unique to this product>"
  }},
  "slide4": {{
    "title": "Your Optimized Listing Preview",
    "new_title_label": "New Title",
    "new_title": "<write a sharp new title for this product under 150 chars>",
    "bullet_1": "<new bullet point 1 starting with BOLD BENEFIT>",
    "bullet_2": "<new bullet point 2 starting with BOLD BENEFIT>",
    "bullet_3": "<new bullet point 3 starting with BOLD BENEFIT>",
    "improvement_note": "<one line: what's fundamentally different about this rewrite>"
  }},
  "slide5": {{
    "title": "Why Work With Us",
    "value_prop": "<one bold value proposition line>",
    "point_1": "<specific thing we do>",
    "point_2": "<specific guarantee or process step>",
    "point_3": "<specific result or social proof>",
    "cta": "<clear call to action>",
    "risk_reversal": "<remove risk>"
  }}
}}
"""

    try:
        return _extract_json(generate_text(prompt, max_tokens=1500))
    except Exception as e:
        if "LLM_PROVIDER=local" not in str(e):
            print(f"Pitch content generation failed, using local fallback: {e}")
        return fallback_pitch_content(giant)


def _set_bg(slide, color):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def _add_text(
    slide,
    text,
    left,
    top,
    width,
    height,
    font_size=18,
    bold=False,
    color=WHITE,
    align=PP_ALIGN.LEFT,
    italic=False,
):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    run = paragraph.add_run()
    run.text = str(text)
    run.font.size = Pt(font_size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.italic = italic
    return box


def _add_rect(slide, left, top, width, height, color):
    shape = slide.shapes.add_shape(1, Inches(left), Inches(top), Inches(width), Inches(height))
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    return shape


def build_slide_1(prs, data):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _set_bg(slide, DARK)
    _add_rect(slide, 0, 0, 10, 0.08, ORANGE)
    _add_rect(slide, 0.3, 0.2, 0.4, 0.35, ORANGE)
    _add_text(slide, "01", 0.3, 0.2, 0.4, 0.35, 12, True, WHITE, PP_ALIGN.CENTER)
    _add_text(slide, data["title"], 0.85, 0.15, 9.0, 0.5, 14, True, ORANGE)
    _add_text(slide, data["product_name"], 0.3, 0.75, 9.4, 1.0, 30, True, WHITE)
    _add_text(
        slide,
        data["tagline"],
        0.3,
        1.6,
        9.4,
        0.5,
        16,
        False,
        RGBColor(0xAA, 0xAA, 0xAA),
        italic=True,
    )
    _add_rect(slide, 0.3, 2.2, 9.4, 0.03, ORANGE)

    metrics = [
        ("MONTHLY REVENUE", data["monthly_revenue"]),
        ("STAR RATING", data["star_rating"]),
        ("REVIEWS", data["review_count"]),
        ("LISTING SCORE", data["listing_score"]),
    ]
    score_val = int(str(data["listing_score"]).replace("/100", "") or 0)
    score_colors = [
        WHITE,
        WHITE,
        WHITE,
        RED if score_val < 35 else YELLOW if score_val < 55 else GREEN,
    ]

    for idx, (label, value) in enumerate(metrics):
        left = 0.3 + idx * 2.4
        _add_rect(slide, left, 2.35, 2.2, 1.15, RGBColor(0x2A, 0x2A, 0x3E))
        _add_text(slide, label, left + 0.12, 2.45, 2.0, 0.3, 8, True, ORANGE)
        _add_text(slide, value, left + 0.12, 2.78, 2.0, 0.55, 20, True, score_colors[idx])

    _add_text(
        slide,
        "CONFIDENTIAL - PREPARED BY SLEEPING GIANT DETECTOR",
        0.3,
        6.9,
        9.4,
        0.3,
        8,
        False,
        RGBColor(0x55, 0x55, 0x55),
        PP_ALIGN.CENTER,
    )


def build_slide_2(prs, data):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _set_bg(slide, DARK)
    _add_rect(slide, 0, 0, 10, 0.08, RED)
    _add_rect(slide, 0.3, 0.2, 0.4, 0.35, RED)
    _add_text(slide, "02", 0.3, 0.2, 0.4, 0.35, 12, True, WHITE, PP_ALIGN.CENTER)
    _add_text(slide, data["title"], 0.85, 0.15, 9.0, 0.5, 14, True, RED)
    _add_text(
        slide,
        data["subtitle"],
        0.3,
        0.75,
        9.4,
        0.5,
        16,
        False,
        RGBColor(0xCC, 0xCC, 0xCC),
        italic=True,
    )
    _add_rect(slide, 0.3, 1.35, 9.4, 0.03, RED)

    weaknesses = [data["weakness_1"], data["weakness_2"], data["weakness_3"]]
    for idx, weakness in enumerate(weaknesses):
        top = 1.55 + idx * 1.4
        _add_rect(slide, 0.3, top, 9.4, 1.2, RGBColor(0x2A, 0x1A, 0x1A))
        _add_rect(slide, 0.3, top, 0.08, 1.2, RED)
        _add_text(slide, str(idx + 1), 0.55, top + 0.35, 0.4, 0.5, 22, True, RED)
        _add_text(slide, weakness, 1.05, top + 0.15, 8.4, 0.9, 13, False, WHITE)

    _add_rect(slide, 0.3, 5.85, 9.4, 0.65, RGBColor(0x3A, 0x0A, 0x0A))
    _add_text(slide, "ALERT: " + data["closing_line"], 0.5, 5.92, 9.0, 0.5, 13, True, YELLOW)


def build_slide_3(prs, data):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _set_bg(slide, DARK)
    _add_rect(slide, 0, 0, 10, 0.08, GREEN)
    _add_rect(slide, 0.3, 0.2, 0.4, 0.35, GREEN)
    _add_text(slide, "03", 0.3, 0.2, 0.4, 0.35, 12, True, WHITE, PP_ALIGN.CENTER)
    _add_text(slide, data["title"], 0.85, 0.15, 9.0, 0.5, 14, True, GREEN)
    _add_rect(slide, 0.3, 0.75, 9.4, 0.03, GREEN)

    _add_rect(slide, 0.3, 0.9, 2.9, 1.7, RGBColor(0x2A, 0x2A, 0x3E))
    _add_rect(slide, 3.55, 0.9, 2.9, 1.7, RGBColor(0x3A, 0x1A, 0x1A))
    _add_rect(slide, 6.8, 0.9, 2.9, 1.7, RGBColor(0x1A, 0x3A, 0x1A))

    _add_text(slide, "CURRENT REVENUE", 0.45, 0.98, 2.6, 0.3, 9, True, RGBColor(0x99, 0x99, 0x99))
    _add_text(slide, data["current_revenue"], 0.45, 1.3, 2.6, 0.9, 22, True, WHITE)
    _add_text(slide, "per month", 0.45, 2.1, 2.6, 0.35, 10, False, RGBColor(0x99, 0x99, 0x99))

    _add_text(slide, "MONEY BEING LOST", 3.7, 0.98, 2.6, 0.3, 9, True, RED)
    _add_text(slide, data["money_left"], 3.7, 1.3, 2.6, 0.9, 22, True, RED)
    _add_text(slide, "being left on table", 3.7, 2.1, 2.6, 0.35, 10, False, RED)

    _add_text(slide, "WITH OPTIMIZATION", 6.95, 0.98, 2.6, 0.3, 9, True, GREEN)
    _add_text(slide, data["optimized_revenue"], 6.95, 1.3, 2.6, 0.9, 22, True, GREEN)
    _add_text(slide, f"+{data['uplift_percent']} potential", 6.95, 2.1, 2.6, 0.35, 10, False, GREEN)

    _add_text(slide, "->", 3.25, 1.45, 0.4, 0.5, 18, True, ORANGE, PP_ALIGN.CENTER)
    _add_text(slide, "->", 6.5, 1.45, 0.4, 0.5, 18, True, ORANGE, PP_ALIGN.CENTER)

    _add_rect(slide, 0.3, 2.85, 4.6, 1.5, RGBColor(0x22, 0x22, 0x35))
    _add_rect(slide, 5.1, 2.85, 4.6, 1.5, RGBColor(0x22, 0x22, 0x35))
    _add_text(slide, data["stat_1"], 0.5, 2.95, 4.2, 1.2, 12, False, WHITE)
    _add_text(slide, data["stat_2"], 5.3, 2.95, 4.2, 1.2, 12, False, WHITE)

    _add_rect(slide, 0.3, 4.55, 9.4, 0.7, ORANGE)
    _add_text(slide, "INSIGHT: " + data["insight"], 0.5, 4.65, 9.0, 0.5, 13, True, DARK)


def build_slide_4(prs, data):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _set_bg(slide, DARK)
    _add_rect(slide, 0, 0, 10, 0.08, ORANGE)
    _add_rect(slide, 0.3, 0.2, 0.4, 0.35, ORANGE)
    _add_text(slide, "04", 0.3, 0.2, 0.4, 0.35, 12, True, WHITE, PP_ALIGN.CENTER)
    _add_text(slide, data["title"], 0.85, 0.15, 9.0, 0.5, 14, True, ORANGE)

    _add_rect(slide, 0.3, 0.75, 9.4, 0.3, ORANGE)
    _add_text(slide, "NEW TITLE", 0.4, 0.77, 1.5, 0.26, 9, True, DARK)
    _add_rect(slide, 0.3, 1.05, 9.4, 0.75, RGBColor(0x1E, 0x2A, 0x1E))
    _add_text(slide, data["new_title"], 0.45, 1.1, 9.1, 0.65, 13, True, GREEN)

    _add_rect(slide, 0.3, 1.95, 9.4, 0.3, ORANGE)
    _add_text(slide, "OPTIMIZED BULLET POINTS", 0.4, 1.97, 4.0, 0.26, 9, True, DARK)

    for idx, bullet in enumerate([data["bullet_1"], data["bullet_2"], data["bullet_3"]]):
        top = 2.35 + idx * 1.1
        _add_rect(slide, 0.3, top, 9.4, 0.95, RGBColor(0x1A, 0x1A, 0x2A))
        _add_rect(slide, 0.3, top, 0.06, 0.95, GREEN)
        _add_text(slide, f"0{idx + 1}", 0.5, top + 0.28, 0.4, 0.4, 13, True, ORANGE)
        _add_text(slide, bullet, 0.95, top + 0.1, 8.5, 0.75, 12, False, WHITE)

    _add_rect(slide, 0.3, 5.75, 9.4, 0.55, RGBColor(0x1E, 0x2A, 0x1E))
    _add_text(slide, "WIN: " + data["improvement_note"], 0.5, 5.83, 9.0, 0.4, 12, True, GREEN)


def build_slide_5(prs, data):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _set_bg(slide, DARK)
    _add_rect(slide, 0, 0, 10, 0.08, ORANGE)
    _add_rect(slide, 0.3, 0.2, 0.4, 0.35, ORANGE)
    _add_text(slide, "05", 0.3, 0.2, 0.4, 0.35, 12, True, WHITE, PP_ALIGN.CENTER)
    _add_text(slide, data["title"], 0.85, 0.15, 9.0, 0.5, 14, True, ORANGE)
    _add_text(slide, data["value_prop"], 0.3, 0.75, 9.4, 0.85, 24, True, WHITE)
    _add_rect(slide, 0.3, 1.65, 9.4, 0.04, ORANGE)

    for idx, point in enumerate([data["point_1"], data["point_2"], data["point_3"]]):
        top = 1.85 + idx * 1.05
        _add_rect(slide, 0.3, top, 9.4, 0.88, RGBColor(0x22, 0x22, 0x35))
        _add_rect(slide, 0.3, top, 0.06, 0.88, ORANGE)
        _add_text(slide, f"{idx + 1}.", 0.5, top + 0.2, 0.5, 0.5, 18, True, ORANGE)
        _add_text(slide, point, 1.1, top + 0.15, 8.4, 0.6, 13, False, WHITE)

    _add_rect(slide, 0.3, 5.1, 9.4, 0.85, ORANGE)
    _add_text(slide, data["cta"], 0.5, 5.18, 9.0, 0.55, 16, True, DARK)
    _add_text(
        slide,
        data["risk_reversal"],
        0.3,
        6.1,
        9.4,
        0.4,
        11,
        False,
        RGBColor(0x99, 0x99, 0x99),
        PP_ALIGN.CENTER,
        italic=True,
    )


def generate_pitch_deck(giant):
    """
    Full pipeline: generate content, build slides, return bytes.
    Returns (pptx_bytes, slide_content_dict).
    """
    content = generate_pitch_content(giant)
    if not content:
        return None, None

    prs = Presentation()
    prs.slide_width = Inches(10)
    prs.slide_height = Inches(7.5)

    build_slide_1(prs, content["slide1"])
    build_slide_2(prs, content["slide2"])
    build_slide_3(prs, content["slide3"])
    build_slide_4(prs, content["slide4"])
    build_slide_5(prs, content["slide5"])

    buf = io.BytesIO()
    prs.save(buf)
    buf.seek(0)

    return buf.getvalue(), content
